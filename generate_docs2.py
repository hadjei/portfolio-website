import sys
import subprocess

# Install python-docx if not available
try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_horizontal_line(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    pPr.insert_element_before(pBdr,
        'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap',
        'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
        'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
        'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
        'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap',
        'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr',
        'w:pPrChange'
    )
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'auto')
    pBdr.append(bottom)

def set_font(doc, font_name="Times New Roman", font_size=10):
    style = doc.styles['Normal']
    font = style.font
    font.name = font_name
    font.size = Pt(font_size)

# Create Resume
doc = Document()
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

set_font(doc, "Times New Roman", 10)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Hafiz Adjei")
run.bold = True
run.font.size = Pt(14)

p = doc.add_paragraph("Accra, Ghana | +233 54 920 1359 | https://hadjei.github.io/portfolio-website/ | adjeihafiz2@gmail.com")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def add_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    add_horizontal_line(p)

def add_entry(doc, org, loc, role, dates, bullets):
    p = doc.add_paragraph()
    run1 = p.add_run(org)
    run1.bold = True
    if loc:
        p.add_run(f"  {loc}")
    
    # tab stop to right align the date
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:pos'), '9360')
    tabs.append(tab)
    pPr.append(tabs)
    
    run_date = p.add_run(f"\t{dates}")
    run_date.bold = True

    if role:
        p_role = doc.add_paragraph()
        run_role = p_role.add_run(role)
        run_role.italic = True
    
    for bullet in bullets:
        p_bullet = doc.add_paragraph(bullet, style='List Bullet')
        p_bullet.paragraph_format.space_after = Pt(2)

add_heading(doc, "EDUCATION")
add_entry(doc, "Carnegie Mellon University Africa", "Kigali, Rwanda", "Master of Science in Electrical and Computer Engineering", "Aug. 2025 - May 2026", [
    "Concentrations: Wireless Communications & Network Applications, Deep Learning, IoT, Applied Machine Learning",
    "Cumulative GPA: 3.69/4.0"
])
add_entry(doc, "Ashesi University", "Berekusu, Ghana", "Bachelor of Science in Electrical and Electronic Engineering", "Sep. 2020 - Aug. 2024", [
    "Honors: Dean's List (2021-2024), MasterCard Foundation Scholar (full merit scholarship)",
    "Cumulative GPA: 3.53/4.0"
])

add_heading(doc, "WORK EXPERIENCE/INTERNSHIPS")
add_entry(doc, "KCRC / CMU Africa", "Kigali, Rwanda", "Research Associate", "May 2026 - Present", [
    "Reverse-engineered a proprietary 5V TTL UART protocol from legacy documentation to integrate battery modules with a modern monitoring system",
    "Developed a full-stack BMS monitoring system comprising a Python serial communication library, Flask REST API, and a real-time web dashboard",
    "Implemented multi-module daisy-chain addressing to poll 3 battery modules simultaneously with validated packets"
])
add_entry(doc, "Carnegie Mellon University Africa", "Kigali, Rwanda", "Student IT Support", "Sep. 2025 - May 2026", [
    "Delivered real-time AV technical support during lectures and large-scale university events, maintaining near-zero downtime",
    "Assisted network and server engineers with LAN administration, fault diagnosis, and data-centre operations"
])
add_entry(doc, "Northern Electricity Distribution Company", "Tamale, Ghana", "Electrical Technician Intern", "Jul. 2023 - Aug. 2023", [
    "Executed service-drop installations and electricity meter deployments for 500+ customers, expanding grid access across northern Ghana",
    "Verified monthly billing data against meter-reader records to ensure billing accuracy and reduce revenue leakage"
])
add_entry(doc, "Think Education", "Tamale, Ghana", "Research Intern", "Jul. 2022 - Aug. 2022", [
    "Designed and populated a performance database for 20+ low-cost private schools, enabling data-driven benchmarking",
    "Analyzed school administration, management, and operations; produced structured reports with actionable improvement recommendations"
])

add_heading(doc, "PROJECTS")
add_entry(doc, "Smart City Parking System", "CMU Africa", "", "Nov. 2025 - Dec. 2025", [
    "Designed a low-cost smart parking system integrating inductive sensing, ultrasonic validation, and edge computing for high-accuracy vehicle detection",
    "Built a web-based dashboard for live slot-occupancy monitoring, real-time analytics, and local session tracking"
])
add_entry(doc, "Autonomous Ground Robot for Greenhouse Monitoring", "Ashesi University", "", "Jan. 2024 - Oct. 2024", [
    "Designed and deployed a real-time environmental monitoring system enabling data-driven crop optimisation for greenhouse farmers",
    "Engineered autonomous navigation and path-planning algorithms for the robot; built PHP/HTML dashboards for remote monitoring"
])

add_heading(doc, "LEADERSHIP EXPERIENCE")
add_entry(doc, "Ashesi Students Council", "Berekusu, Ghana", "Academic Committee Member", "Mar. 2023 - Apr. 2024", [
    "Supervised peer-tutoring programme covering mathematics and computer programming for underclassmen",
    "Spearheaded academic integrity awareness campaigns to deter plagiarism campus-wide"
])

add_heading(doc, "CORE COMPETENCIES AND SKILLS")
doc.add_paragraph("Data Analytics & Programming: Python, C, JavaScript, PHP, HTML, Data Analytics, TensorFlow, Deep Learning", style='List Bullet')
doc.add_paragraph("Consulting Tools: Microsoft Excel (Advanced), Microsoft PowerPoint (Advanced), Structured Reporting", style='List Bullet')
doc.add_paragraph("Engineering & IoT: Embedded Systems, ARM Assembly, ESP32, Arduino, MQTT, Bluetooth, Technical Troubleshooting", style='List Bullet')

add_heading(doc, "CERTIFICATIONS")
doc.add_paragraph("Harvard Health Systems Innovation Lab Hackathon: Certificate in Building High-Value Health Systems Leveraging AI (Apr. 2025)", style='List Bullet')

doc.save("roland_berger_resume_with_projects.docx")
print("Successfully generated roland_berger_resume_with_projects.docx with Projects")
